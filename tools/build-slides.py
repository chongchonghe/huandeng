#!/usr/bin/env python3
"""Build and check a Quarto + reveal.js deck.

    build-slides.py <deck>                  HTML, PDF and PPTX
    build-slides.py <deck> --html
    build-slides.py <deck> --standalone     one self-contained .html to email
    build-slides.py <deck> --pdf
    build-slides.py <deck> --pptx           one slide image per page, at 4K
    build-slides.py <deck> --check          render, then look at every slide
    build-slides.py <deck> --png            one PNG per slide, to read
    build-slides.py <deck> --gallery        every theme side by side, to choose one

`--html` and `--standalone` are `quarto render` with the right flags, and you
can type those yourself. `--pdf`, `--png` and `--check` drive a headless
Chromium through Playwright, because they need a browser to have laid the slides
out before there is anything to measure or print; `--pptx` then rasterises that
PDF, because a deck whose layout is CSS cannot survive Pandoc's PowerPoint
writer — see `build_pptx`.

`--check` is the reason this file exists. A reveal.js slide that holds too much
does not error and does not shrink: the surplus hangs into the thin margin
reveal keeps around the slide box and is then cut off by the window edge. How
much survives depends on the window's aspect ratio — a 16:10 laptop shows about
57 slide-px of it, a 16:9 projector only 15 — so a slide can look merely tight
while you write it and be cut on stage. Printed, the same slide splits across
two PDF pages instead, title on one and body on the next. The same goes for a
figure given both a width and a height, which silently changes the aspect ratio
of a scientific plot. All of it is caught here by asking the browser where
things actually landed.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
import zipfile
from html import escape
from pathlib import Path

TOLERANCE_PX = 2.0  # slide-space pixels of overflow to forgive
ASPECT_TOLERANCE = 0.02  # 2%, the same figure the Typst side uses


# ---------------------------------------------------------------- the deck --


class Deck:
    def __init__(self, path: Path):
        self.dir = path.resolve()
        if not self.dir.is_dir():
            die(f"no such deck: {path}")
        qmds = sorted(p for p in self.dir.glob("*.qmd") if not p.name.startswith("_"))
        if not qmds:
            die(f"no .qmd in {self.dir}")
        if len(qmds) > 1:
            die(f"more than one .qmd in {self.dir}: {', '.join(q.name for q in qmds)}")
        self.qmd = qmds[0]
        self.name = self.qmd.stem
        self.out = self.dir / "out"
        # Whether out/<name>.html was built during this run. Anything that reads
        # the built deck renders first if it was not: checking a stale HTML is
        # worse than not checking at all, because it comes back green.
        self.fresh = False

    @property
    def html(self) -> Path:
        return self.out / f"{self.name}.html"

    @property
    def standalone(self) -> Path:
        return self.out / f"{self.name}-standalone.html"

    @property
    def pdf(self) -> Path:
        return self.out / f"{self.name}.pdf"


def die(msg: str) -> None:
    print(f"error: {msg}", file=sys.stderr)
    raise SystemExit(1)


# ------------------------------------------------------------------ themes --


# A theme is one stylesheet in the deck's own `themes/`, and every deck carries
# all of them — so changing the look is one line of `_quarto.yml` and nothing
# here has to run. The code-token colours are in the stylesheet too, which is
# why there is no second setting to keep in step with the first.


def themes(deck: Deck) -> list[str]:
    d = deck.dir / "themes"
    return sorted(p.stem for p in d.glob("*.scss")) if d.is_dir() else []


# ------------------------------------------------------------------ quarto --


def quarto(deck: Deck, *args: str) -> None:
    """Render, quietly.

    `quarto render` echoes the whole resolved configuration on every run, which
    buries four decks' worth of `make check` output. `--quiet` suppresses that —
    but it suppresses the error message too, leaving only an exit code. So run
    it quiet, and if it fails, run it again loudly to find out why.
    """
    exe = shutil.which("quarto")
    if exe is None:
        die("quarto is not on PATH — see https://quarto.org/docs/get-started/")
    cmd = [exe, "render", deck.qmd.name, *args]
    print("$ quarto " + " ".join(cmd[1:]), flush=True)
    if subprocess.run([*cmd, "--quiet"], cwd=deck.dir).returncode == 0:
        return
    result = subprocess.run(cmd, cwd=deck.dir)
    raise SystemExit(result.returncode or 1)


def build_html(deck: Deck) -> Path:
    quarto(deck, "--to", "revealjs")
    deck.fresh = True
    return deck.html


def ensure_html(deck: Deck) -> None:
    if not deck.fresh:
        build_html(deck)


def build_standalone(deck: Deck) -> Path:
    """One file with the fonts, the images and MathJax inlined.

    Everything a normal render leaves beside the HTML ends up inside it, so the
    deck can be emailed as a single attachment and opened offline. Videos are
    inlined too, which is why it is a separate target and not the default.
    """
    quarto(
        deck,
        "--to",
        "revealjs",
        "-M",
        "embed-resources:true",
        "--output",
        deck.standalone.name,
    )
    return deck.standalone


def build_pptx(deck: Deck, width_px: int, dpi: int | None) -> Path:
    """PowerPoint, one full-bleed slide image per page — as the Typst decks do.

    Pandoc's own PPTX writer re-flows the Markdown into PowerPoint's layouts,
    and everything that makes a slide a slide here is CSS: the columns, the
    figure credits, the callouts, the theme. What comes out is a bulleted
    outline wearing none of the deck's design. Styling it through a reference
    document reaches the fonts and the colours and no further.

    So do what the Typst side does instead and ship pictures. Each page of the
    PDF — already printed from the real deck in a real browser — becomes one
    image filling one slide. Nothing is editable in PowerPoint, and that is the
    honest trade: PPTX is a delivery format, and this way it delivers the deck
    you actually wrote.

    Pagination follows `<deck>.pdf`, so a build clicks through step by step.

    Sharpness is set as a pixel width, not a print resolution: these are pictures
    of a screen, and what matters is whether there is a source pixel behind every
    display pixel. The default matches a 4K projector across the slide's full
    width. Anything less is upscaled on the night and looks soft.
    """
    import pymupdf
    from pptx import Presentation
    from pptx.util import Inches

    pdf = deck.pdf
    if not pdf.exists() or not deck.fresh:
        build_pdf(deck)

    doc = pymupdf.open(pdf)
    page_w_in = doc[0].rect.width / 72
    # `--dpi` still wins if someone asks for one by name, so the flag means the
    # same thing here as it does on the Typst side.
    render_dpi = dpi if dpi is not None else round(width_px / page_w_in)

    pres = Presentation()
    # The page is already the slide's shape; take it verbatim rather than
    # rounding to PowerPoint's nominal widescreen, which would rescale every
    # image by a hair.
    pres.slide_width = Inches(page_w_in)
    pres.slide_height = Inches(doc[0].rect.height / 72)
    blank = pres.slide_layouts[6]

    px = (0, 0)
    with tempfile.TemporaryDirectory() as scratch:
        for i, page in enumerate(doc):
            frame = Path(scratch) / f"{i + 1:04d}.png"
            pixmap = page.get_pixmap(dpi=render_dpi)
            px = (pixmap.width, pixmap.height)
            pixmap.save(frame)
            slide = pres.slides.add_slide(blank)
            slide.shapes.add_picture(
                str(frame), 0, 0, width=pres.slide_width, height=pres.slide_height
            )

    out = deck.out / f"{deck.name}.pptx"
    pres.save(out)
    normalize_zip(out)
    size_mb = out.stat().st_size / 1e6
    print(f"wrote {out}")
    print(
        f"  {doc.page_count} slides at {px[0]}x{px[1]}px "
        f"({render_dpi} dpi), {size_mb:.1f} MB"
    )
    doc.close()
    return out


ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)


def normalize_zip(path: Path) -> None:
    """Rewrite an archive with fixed timestamps, so identical input gives identical bytes."""
    with zipfile.ZipFile(path) as src:
        members = [(info, src.read(info.filename)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in members:
            fixed = zipfile.ZipInfo(info.filename, date_time=ZIP_EPOCH)
            fixed.compress_type = info.compress_type
            fixed.external_attr = info.external_attr
            out.writestr(fixed, data)


# ----------------------------------------------------------------- browser --


def browser_page(playwright):
    browser = playwright.chromium.launch()
    page = browser.new_page(viewport={"width": 1600, "height": 1000})
    return browser, page


def open_deck(page, html: Path, query: str = "") -> None:
    page.goto(f"file://{html.resolve()}{query}")
    page.wait_for_function("() => window.Reveal && Reveal.isReady()", timeout=30000)
    page.wait_for_timeout(400)


def with_playwright():
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        die(
            "this target needs Playwright:\n"
            "    uv sync --extra quarto\n"
            "    uv run --extra quarto playwright install chromium"
        )
    return sync_playwright()


# Reveal's slide selector is `.slides section` — a *descendant* selector, so any
# <section> anywhere inside a slide silently becomes a slide of its own, with a
# blank entry in the progress bar and an arrow press that goes nowhere.
#
# Markdown gets you there without asking: Pandoc writes a fenced div whose first
# block is a heading as a <section> carrying the div's class, so
# `::: {.card}` + `#### Setup` compiles to `<section class="card">`. Quarto
# rewrites `.column` itself and is safe; nothing else is.
#
# A section reveal means is a child of `.slides` or of another section. Anything
# else got there by accident.
STRAY_SECTION_JS = r"""
() => {
  const out = [];
  document.querySelectorAll('.reveal .slides section').forEach(sec => {
    const p = sec.parentElement;
    if (p.classList.contains('slides') || p.tagName === 'SECTION') return;
    const slide = p.closest('section');
    const all = Reveal.getSlides();
    const h = slide ? slide.querySelector('h1, h2') : null;
    out.push({
      cls: (sec.className || '(no class)').replace(/ ?(past|present|future)\b/g, ''),
      slide: slide ? all.indexOf(slide) + 1 : 0,
      title: h ? h.textContent.trim().slice(0, 40) : 'untitled',
    });
  });
  return out;
}
"""


# Runs inside the page, once per slide, after navigating to it.
#
# Everything is converted back into slide coordinates by dividing out
# Reveal.getScale(), so a measurement means the same thing whatever size the
# window happens to be.
MEASURE_JS = r"""
() => {
  const scale = Reveal.getScale();
  const box = document.querySelector('.reveal .slides').getBoundingClientRect();
  const slide = Reveal.getCurrentSlide();
  // MathJax lays equations out with spans whose boxes reach well above and
  // below the ink — stretchy delimiters and fraction bars are positioned by
  // negative offsets. Measuring those reports overflow that nobody can see, so
  // take the equation's outer container and none of its parts.
  const MATH = '.MathJax, .MathJax_Display, .MathJax_SVG, .MathJax_SVG_Display,'
             + ' .MathJax_Preview, mjx-container';
  // An ancestor that clips its overflow has already decided what is visible, so
  // measuring a child that pokes out of it would report something nobody can
  // see. A numbered code block is the everyday case: Pandoc puts the line
  // numbers in a gutter by pushing each line 4em to the left inside a `pre`
  // that scrolls.
  const clipped = (el) => {
    for (let p = el.parentElement; p && p !== slide; p = p.parentElement) {
      const cs = getComputedStyle(p);
      if (cs.overflowX !== 'visible' || cs.overflowY !== 'visible') return true;
    }
    return false;
  };

  const skip = (el) => {
    const cs = getComputedStyle(el);
    if (cs.display === 'none' || cs.visibility === 'hidden') return true;
    if (el.closest('aside.notes')) return true;
    if (el.classList.contains('MJX_Assistive_MathML')) return true;
    const math = el.closest(MATH);
    if (math && math !== el) return true;
    if (clipped(el)) return true;
    return false;
  };

  let bottom = -Infinity, top = Infinity, right = -Infinity, left = Infinity;
  for (const el of slide.querySelectorAll('*')) {
    if (skip(el)) continue;
    const r = el.getBoundingClientRect();
    if (r.width < 1 || r.height < 1) continue;
    bottom = Math.max(bottom, r.bottom);
    top = Math.min(top, r.top);
    right = Math.max(right, r.right);
    left = Math.min(left, r.left);
  }
  const empty = bottom === -Infinity;

  const media = [];
  for (const el of slide.querySelectorAll('img, video')) {
    if (skip(el)) continue;
    const r = el.getBoundingClientRect();
    const isVideo = el.tagName === 'VIDEO';
    media.push({
      kind: isVideo ? 'video' : 'image',
      src: el.getAttribute('src') || el.getAttribute('data-src') || '(inline)',
      naturalW: isVideo ? el.videoWidth : el.naturalWidth,
      naturalH: isVideo ? el.videoHeight : el.naturalHeight,
      drawnW: r.width / scale,
      drawnH: r.height / scale,
    });
  }

  const heading = slide.querySelector('h1, h2, h3');
  return {
    title: heading ? heading.textContent.trim() : '(no heading)',
    below: empty ? 0 : (bottom - box.bottom) / scale,
    above: empty ? 0 : (box.top - top) / scale,
    past_right: empty ? 0 : (right - box.right) / scale,
    past_left: empty ? 0 : (box.left - left) / scale,
    height: box.height / scale,
    width: box.width / scale,
    media,
  };
}
"""


def visit_slides(page, on_slide) -> int:
    """Walk every slide in order, calling `on_slide(index, measurement)`."""
    total = page.evaluate("() => Reveal.getTotalSlides()")
    indices = page.evaluate(
        "() => Reveal.getSlides().map(s => { const i = Reveal.getIndices(s);"
        " return [i.h, i.v || 0]; })"
    )
    for i, (h, v) in enumerate(indices):
        page.evaluate(f"() => Reveal.slide({h}, {v})")
        # Reveal loads a slide's images the moment it becomes current; give
        # them a chance to arrive before measuring anything against them.
        page.wait_for_function(
            "() => Array.from(Reveal.getCurrentSlide().querySelectorAll('img'))"
            ".every(im => im.complete)",
            timeout=15000,
        )
        page.wait_for_timeout(120)
        on_slide(i, page.evaluate(MEASURE_JS))
    return total


# ------------------------------------------------------------------ checks --


def check(deck: Deck, expect: int | None) -> int:
    ensure_html(deck)

    problems: list[str] = []
    total = 0

    def look(i: int, m: dict) -> None:
        label = f"slide {i + 1} ({m['title']})"
        edges = (
            ("below", "past the bottom", m["height"]),
            ("above", "above the top", m["height"]),
            ("past_right", "past the right edge", m["width"]),
            ("past_left", "past the left edge", m["width"]),
        )
        for key, where, extent in edges:
            if m[key] > TOLERANCE_PX:
                problems.append(
                    f"{label}: content runs {m[key]:.0f}px {where} "
                    f"of a {extent:.0f}px slide"
                )
        for item in m["media"]:
            name = item["src"].rsplit("/", 1)[-1]
            if not item["naturalW"] or not item["naturalH"]:
                if item["kind"] == "image":
                    problems.append(f"{label}: image did not load — {name}")
                continue
            if not item["drawnW"] or not item["drawnH"]:
                continue
            want = item["naturalW"] / item["naturalH"]
            got = item["drawnW"] / item["drawnH"]
            if abs(got - want) / want > ASPECT_TOLERANCE:
                problems.append(
                    f"{label}: {item['kind']} drawn at {got:.3f}:1 but is "
                    f"{want:.3f}:1 — {name}. Give it a width or a height, "
                    f"never both."
                )

    with with_playwright() as pw:
        browser, page = browser_page(pw)
        open_deck(page, deck.html)
        strays = page.evaluate(STRAY_SECTION_JS)
        total = visit_slides(page, look)
        browser.close()

    for s in strays:
        problems.append(
            f"stray <section class=\"{s['cls']}\"> inside slide {s['slide']} "
            f"({s['title']}) — reveal counts it as a slide of its own. A fenced "
            f"div whose first block is a heading is written out as a section; "
            f"lead it with anything else."
        )

    print(f"{deck.name}: {total} slides")
    if expect is not None and total != expect:
        problems.append(f"expected {expect} slides, built {total}")

    if problems:
        print()
        for p in problems:
            print(f"  {p}")
        print(f"\n{len(problems)} problem(s).")
        return 1
    print("no overflow, no stretched figures.")
    return 0


# ------------------------------------------------------------ pdf and png --


# Reveal's print layout wraps each slide in a `.pdf-page` whose height is a
# whole number of printed pages. A slide that holds too much comes out two pages
# tall — the title alone on one page and the body on the next — which is the
# same silent split the Typst decks get, and the same thing a page count catches.
PAGINATION_JS = r"""
() => {
  const pages = Array.from(document.querySelectorAll('.pdf-page'));
  const unit = Math.min(...pages.map(p => p.getBoundingClientRect().height));
  return {
    total: pages.length,
    split: pages.map((p, i) => {
      const n = Math.round(p.getBoundingClientRect().height / unit);
      const h = p.querySelector('h1, h2, h3');
      return n > 1
        ? { i: i + 1, pages: n, title: h ? h.textContent.trim() : '(no heading)' }
        : null;
    }).filter(Boolean),
  };
}
"""


# Two PDFs, from one render. Reveal reads config overrides off the query string,
# so `pdfSeparateFragments` can be flipped per print without re-rendering the
# deck — the same HTML gives both files.
#
#   <deck>.pdf                      one page per *step*: the deck as presented,
#                                   so a build arrives a piece at a time and an
#                                   image sequence becomes a flip-book
#   <deck>-one-page-per-slide.pdf   one page per slide, fully built — the
#                                   handout, and what you read to check a deck
PDF_VARIANTS = (
    ("", "?print-pdf&pdfSeparateFragments=true", "one page per step"),
    ("-one-page-per-slide", "?print-pdf", "one page per slide"),
)


def print_one_pdf(page, html: Path, query: str, path: Path) -> dict:
    open_deck(page, html, query)
    # A video's play bar is furniture for the live deck; on paper it is a grey
    # slab across the poster frame. (The frame is all a PDF can hold — a movie
    # prints as its first frame and there is nothing to be done about that.)
    page.add_style_tag(
        content="video::-webkit-media-controls { display: none !important }"
    )
    page.wait_for_timeout(1200)
    layout = page.evaluate(PAGINATION_JS)
    page.emulate_media(media="print")
    page.pdf(path=str(path), prefer_css_page_size=True, print_background=True)
    return layout


def build_pdf(deck: Deck) -> Path:
    """Print through reveal's own `?print-pdf` layout, twice.

    That mode lays every slide out as a static page of exactly the configured
    slide size, which is what makes a browser print come out one slide per page
    instead of one long scroll.
    """
    ensure_html(deck)
    with with_playwright() as pw:
        browser, page = browser_page(pw)
        for suffix, query, what in PDF_VARIANTS:
            path = deck.out / f"{deck.name}{suffix}.pdf"
            layout = print_one_pdf(page, deck.html, query, path)
            print(f"wrote {path}")
            if layout["split"]:
                extra = sum(s["pages"] - 1 for s in layout["split"])
                print(f"  {layout['total'] + extra} pages, {what} — except:")
                for s in layout["split"]:
                    print(f"    {s['i']} ({s['title']}) spills onto {s['pages']} pages")
                print("  run --check to see how far past the edge each one goes")
            else:
                print(f"  {layout['total']} pages, {what}")
        browser.close()
    return deck.pdf


def build_png(deck: Deck) -> Path:
    """One PNG per slide. 'It rendered' is not evidence that a slide is legible."""
    ensure_html(deck)
    png_dir = deck.out / "png"
    png_dir.mkdir(parents=True, exist_ok=True)
    for old in png_dir.glob("slide-*.png"):
        old.unlink()

    with with_playwright() as pw:
        browser = pw.chromium.launch()
        page = browser.new_page(viewport={"width": 1280, "height": 720})
        open_deck(page, deck.html)
        # A viewport exactly one slide tall makes each screenshot the slide and
        # nothing else — no letterboxing to crop off afterwards.
        page.add_style_tag(
            content=".reveal .progress, .reveal .slide-menu-button { display: none }"
        )
        written = []

        def shoot(i: int, _m: dict) -> None:
            path = png_dir / f"slide-{i + 1:03d}.png"
            page.screenshot(path=str(path))
            written.append(path)

        visit_slides(page, shoot)
        browser.close()
    print(f"wrote {len(written)} PNGs to {png_dir}")
    return png_dir


# ----------------------------------------------------------------- gallery --


def theme_blurb(deck: Deck, name: str) -> str:
    """The theme's own opening sentence, read from the top of its stylesheet.

    Taken from the file rather than written out here, because a second copy of
    every theme's description is a second thing that has to stay true.
    `themes/README.md` holds the long form; this is the line the theme leads
    with, and it moves when the theme does.
    """
    blurb: list[str] = []
    for line in (deck.dir / "themes" / f"{name}.scss").read_text().splitlines():
        text = line.strip()
        body = text[2:].strip() if text.startswith("//") else ""
        if not blurb:
            if body:
                blurb.append(body)
            continue
        if not body:  # the blank `//` that closes the opening paragraph
            break
        blurb.append(body)
    out = " ".join(blurb)
    for dash in (" — ", " - "):
        if out.lower().startswith(name.lower() + dash):
            return out[len(name) + len(dash) :]
    return out


# Which slides of a theme's template deck earn a column. Four is enough to tell
# them apart, and a shorter row is a row you can take in at once: the title
# slide for the ground and the type, then the three that carry the most theme —
# bullets and inline code, the type scale with a table and a highlighted code
# block, and the cards. Numbered as the deck numbers them, so a file called
# swiss-05.png is slide 5 of swiss and nothing has to be counted back.
GALLERY_SLIDES = (1, 3, 5, 6)

# The order the gallery shows them in, grouped by what the slide is printed on.
# Ground first, because it is the first thing that has to suit your figures: a
# plot saved on a white canvas is a bright rectangle punched into a dark slide,
# and no stylesheet fixes that. Inside a group the order is the author's, not
# alphabetical — none of this can be derived from the files, which is why it is
# written down rather than computed.
#
# A theme missing from this list still appears, at the end. The gallery never
# silently drops one; the deck's own `themes/README.md` stays alphabetical.
GALLERY_GROUPS = (
    ("white", ("swiss", "university", "whiteprint", "cobalt", "monochrome")),
    ("light colour", ("paper", "signal", "solarized")),
    ("dark", ("blueprint", "nord")),
)


def gallery_order(deck: Deck) -> list[tuple[str, str]]:
    """Every theme as (group, name), in the order the gallery lays them out."""
    have = themes(deck)
    out = [(g, n) for g, names in GALLERY_GROUPS for n in names if n in have]
    placed = {n for _, n in out}
    return out + [("not yet placed", n) for n in have if n not in placed]


GALLERY_CSS = """
:root { color-scheme: light dark; --ink: #14161a; --dim: #5d6470;
        --ground: #f4f5f7; --card: #fff; --line: #d8dbe0; }
@media (prefers-color-scheme: dark) {
  :root { --ink: #e8eaed; --dim: #99a0ab; --ground: #16181c;
          --card: #1e2126; --line: #333840; }
}
* { box-sizing: border-box; }
body { margin: 0; padding: 2rem clamp(1rem, 4vw, 3rem); background: var(--ground);
       color: var(--ink); font: 15px/1.5 -apple-system, "Fira Sans", Segoe UI, sans-serif; }
h1 { font-size: 1.5rem; margin: 0 0 .3rem; }
p.lede { color: var(--dim); margin: 0 0 .6rem; max-width: 78ch; }
p.lede:last-of-type { margin-bottom: 1.8rem; }
code { font-family: ui-monospace, SFMono-Regular, Menlo, monospace; font-size: .92em; }

/* The grid is the point: every theme deck holds the same slides, so a column
   is one slide in every theme and a row is one theme end to end. It scrolls
   sideways rather than reflowing, because a row that wrapped would stop
   lining up with the row above it. */
.sheet { overflow-x: auto; background: var(--card); border: 1px solid var(--line);
         border-radius: 8px; }
/* Four columns share whatever is left over, so the sheet fits an ordinary
   window and the slides get as big as that allows. Below the min-width the
   container scrolls rather than shrinking them past reading. */
table { border-collapse: collapse; width: 100%; table-layout: fixed;
        min-width: 60rem; }
th, td { padding: 0; vertical-align: top; }
th.slide-head { font-weight: 600; font-size: .78rem; text-align: left;
                padding: .55rem .7rem; color: var(--dim); white-space: nowrap;
                border-bottom: 1px solid var(--line); }
th.slide-head b { color: var(--ink); font-weight: 700; }
/* Sticky, so the name stays with the pictures however far right you scroll. */
th.theme, th.corner { position: sticky; left: 0; z-index: 2; background: var(--card);
                      border-right: 1px solid var(--line); text-align: left;
                      padding: .8rem .9rem; width: 17rem; min-width: 17rem; }
th.theme { border-top: 1px solid var(--line); vertical-align: top; }
th.theme .name { font-family: ui-monospace, SFMono-Regular, Menlo, monospace;
                 font-size: .95rem; font-weight: 700; }
th.theme .blurb { font-weight: 400; font-size: .8rem; color: var(--dim);
                  margin-top: .35rem; }
th.theme .cmd { font-size: .62rem; color: var(--dim); margin-top: .6rem;
                line-height: 1.5; overflow-wrap: break-word;
                font-family: ui-monospace, SFMono-Regular, Menlo, monospace; }
/* The band that says what the next few rows are printed on. */
th.group { background: var(--ground); border-top: 1px solid var(--line);
           color: var(--dim); font-size: .72rem; font-weight: 600;
           letter-spacing: .14em; padding: .5rem .9rem; text-align: left;
           text-transform: uppercase; }
td.shot { border-top: 1px solid var(--line); }
/* Width only, never a height as well: these are 1280x720 screenshots and a
   second dimension would squash them off ratio. */
td.shot a { display: block; }
td.shot img { width: 100%; height: auto; display: block; }
td.shot a:hover img { outline: 2px solid currentColor; outline-offset: -2px; }
</style>
"""


def gallery_html(
    deck: Deck, order: list[tuple[str, str]], titles: list[str],
    shots: dict[str, list[Path]]
) -> str:
    head = "".join(
        f'<th class="slide-head"><b>{GALLERY_SLIDES[i]}</b> {escape(t)}</th>'
        for i, t in enumerate(titles)
    )
    rows = []
    seen_group = None
    for group, name in order:
        if group != seen_group:
            seen_group = group
            rows.append(
                f'<tr><th class="group" colspan="{len(titles) + 1}">'
                f"{escape(group)}</th></tr>"
            )
        # Each thumbnail links to its own PNG, which is the whole 1280x720
        # slide: small enough to compare every theme at a glance, one click
        # from big enough to read.
        cells = "".join(
            f'<td class="shot"><a href="gallery/{p.name}">'
            f'<img src="gallery/{p.name}" alt="{escape(name)}, slide {GALLERY_SLIDES[i]}">'
            f"</a></td>"
            for i, p in enumerate(shots[name])
        )
        rows.append(
            f'<tr><th class="theme"><div class="name">{escape(name)}</div>'
            f'<div class="blurb">{escape(theme_blurb(deck, name))}</div>'
            f'<div class="cmd">theme: [default, themes/{escape(name)}.scss]</div>'
            f"</th>{cells}</tr>"
        )
    return (
        '<!doctype html><meta charset="utf-8">\n'
        "<title>huandeng themes</title>\n<style>"
        + GALLERY_CSS
        + "\n<h1>The starting points</h1>\n"
        '<p class="lede">Every theme this deck can wear, on the four slides that show '
        "the most of it. Read a row for one theme, or a column to compare the same slide "
        "across all of them. Any slide opens full size if you click it.</p>\n"
        '<p class="lede">They are grouped by what the slide is printed on, because that '
        "is the first thing that has to suit your figures: a plot saved on a white "
        "canvas is a bright rectangle punched into a dark slide, and no stylesheet "
        "fixes that.</p>\n"
        '<p class="lede">Pick one and put its line into this deck\'s '
        "<code>_quarto.yml</code> — it is printed under each name. Every stylesheet "
        "already travels with the deck, so nothing is copied and undo is undo.</p>\n"
        f'<div class="sheet"><table><tr><th class="corner"></th>{head}</tr>\n'
        + "\n".join(rows)
        + "</table></div>\n"
    )


def dress(work: Deck, name: str) -> None:
    """Point a working copy's `_quarto.yml` at one of its own stylesheets.

    The same one line a person edits by hand, edited by hand — there is no
    theme machinery left for this to call, which is the point of the design.
    """
    yml = work.dir / "_quarto.yml"
    out = []
    for line in yml.read_text().splitlines(keepends=True):
        head = line.lstrip()
        if head.startswith("theme:"):
            pad = line[: len(line) - len(head)]
            line = f"{pad}theme: [default, themes/{name}.scss]\n"
        out.append(line)
    yml.write_text("".join(out))


def build_gallery(deck: Deck) -> Path:
    """Every theme, side by side, so a person can point at one.

    A name is not something anyone can choose between, and this deck carries
    every theme it could wear — so the gallery is this deck rendered once per
    stylesheet. Every row is the same slides in a different look, which is what
    makes reading down a column mean anything.

    They are laid out grouped by ground, in `GALLERY_GROUPS` — not
    alphabetically, because the order is a judgement about what to look at
    first and nothing in the files can tell you that.

    Your deck is never touched: the dressing happens to a copy under /tmp, so
    running this halfway through writing a talk cannot cost you the look you
    already chose. The sheet lands in your deck's own `out/`.
    """
    order = gallery_order(deck)
    if not order:
        die(f"no themes in {deck.dir / 'themes'}")
    names = [n for _, n in order]

    gallery = deck.out / "gallery"
    gallery.mkdir(parents=True, exist_ok=True)
    for old in gallery.glob("*.png"):
        old.unlink()

    shots: dict[str, list[Path]] = {}
    titles: list[str] = []

    with tempfile.TemporaryDirectory() as tmp:
        work = Deck(
            Path(
                shutil.copytree(
                    deck.dir,
                    Path(tmp) / deck.dir.name,
                    ignore=shutil.ignore_patterns("out", ".quarto"),
                )
            )
        )
        with with_playwright() as pw:
            browser = pw.chromium.launch()
            page = browser.new_page(viewport={"width": 1280, "height": 720})
            for name in names:
                print(f"=== {name}", flush=True)
                dress(work, name)
                work.fresh = False
                build_html(work)
                open_deck(page, work.html)
                # The progress bar stays. `make png` hides it as furniture in
                # the way of reading a slide; here it is one of the things
                # being chosen between, since three themes do not draw one.
                written: list[Path] = []

                def shoot(i: int, m: dict, _name: str = name) -> None:
                    if i + 1 not in GALLERY_SLIDES:
                        return
                    path = gallery / f"{_name}-{i + 1:02d}.png"
                    page.screenshot(path=str(path))
                    written.append(path)
                    if _name == names[0]:
                        titles.append(m["title"])

                visit_slides(page, shoot)
                shots[name] = written
            browser.close()

    sheet = deck.out / "gallery.html"
    sheet.write_text(gallery_html(deck, order, titles, shots))
    n = sum(len(v) for v in shots.values())
    print(f"\nwrote {sheet}")
    print(f"  {len(names)} themes, {n} slides. Open it and point at one.")
    return sheet


# -------------------------------------------------------------------- main --


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("deck", type=Path, help="the deck directory")
    ap.add_argument("--html", action="store_true")
    ap.add_argument(
        "--standalone", action="store_true", help="one self-contained .html file"
    )
    ap.add_argument("--pdf", action="store_true")
    ap.add_argument("--pptx", action="store_true")
    ap.add_argument(
        "--width",
        type=int,
        default=3840,
        metavar="PX",
        help="PPTX slide image width in pixels (default 3840, a 4K projector)",
    )
    ap.add_argument(
        "--dpi", type=int, help="PPTX slide image DPI, overriding --width"
    )
    ap.add_argument("--png", action="store_true", help="one PNG per slide")
    ap.add_argument(
        "--gallery",
        action="store_true",
        help="every theme, side by side, as out/gallery.html",
    )
    ap.add_argument("--check", action="store_true", help="overflow and aspect ratio")
    ap.add_argument(
        "--expect", type=int, metavar="N", help="fail unless the deck has N slides"
    )
    args = ap.parse_args(argv)

    deck = Deck(args.deck)

    # The gallery is about themes/ rather than about this deck — the deck only
    # says where to put it — so it does the one thing and builds nothing else.
    if args.gallery:
        build_gallery(deck)
        return 0

    picked = any([args.html, args.standalone, args.pdf, args.pptx, args.png, args.check])

    if not picked:
        args.html = args.pdf = args.pptx = True

    if args.html:
        build_html(deck)
    if args.standalone:
        build_standalone(deck)
    if args.pdf:
        build_pdf(deck)
    if args.pptx:
        build_pptx(deck, args.width, args.dpi)
    if args.png:
        build_png(deck)
    # Last, so its verdict is the last thing on screen and the exit code.
    if args.check:
        return check(deck, args.expect)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
